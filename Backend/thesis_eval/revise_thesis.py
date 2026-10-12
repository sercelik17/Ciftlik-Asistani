# -*- coding: utf-8 -*-
"""
Tez Word dosyasını 10 Ekim 2026 tarihli yeniden değerlendirme raporuna göre düzenler.

Önce veri paketini üretin:
    cd Backend
    python -m thesis_eval.build_thesis_revision

Sonra tezi düzenleyin:
    python -m thesis_eval.revise_thesis --src "C:/Users/Serenay/Downloads/hayvancılık güncel son.docx"

İsteğe bağlı:
    --dst  <çıktı.docx>          (varsayılan: <src>_revize_2026-10-10.docx)
    --retitle                    kapak/başlık paragraflarını süt sığırcılığı başlığıyla değiştirir
    --no-appendices              Ek Tablo 2–4 ve prompt ekini basma
    --dry-run                    dosyayı kaydetmeden hangi işlemlerin uygulanacağını raporla

Betik, bulamadığı çapaları (anchor) sonunda "UYGULANAMADI" olarak listeler; bunları
evaluation_output/revizyon_2026_10_10/tez_revizyon_paketi.html içinden elle yapıştırın.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

BACKEND_DIR = Path(__file__).resolve().parents[1]
DEFAULT_DATA = BACKEND_DIR / "evaluation_output" / "revizyon_2026_10_10" / "revision_data.json"

OLD_TITLE_RE = re.compile(
    r"hayvanc[ıi]l[ıi]k\s+sekt[öo]r[üu]nde\s+yapay\s+zek[âa]\s+destekli\s+sohbet\s+robotu",
    re.IGNORECASE,
)

LOG: list[tuple[str, str, str]] = []


def log(op: str, status: str, detail: str = "") -> None:
    LOG.append((op, status, detail))


# ----------------------------------------------------------------------------
# docx yardımcıları
# ----------------------------------------------------------------------------
def set_para_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_paragraph_after(paragraph: Paragraph, text: str, style=None, bold: bool = False) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        try:
            para.style = style
        except Exception:
            pass
    run = para.add_run(text)
    if bold:
        run.bold = True
    return para


def insert_paragraph_before(paragraph: Paragraph, text: str, style=None, bold: bool = False) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        try:
            para.style = style
        except Exception:
            pass
    run = para.add_run(text)
    if bold:
        run.bold = True
    return para


def insert_table_after(doc: Document, paragraph: Paragraph, headers: list[str], rows: list[list], table_style=None):
    table = doc.add_table(rows=1, cols=len(headers))
    if table_style is not None:
        try:
            table.style = table_style
        except Exception:
            pass
    for j, head in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(head))
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = "" if value is None else str(value)
    paragraph._p.addnext(table._tbl)
    return table


def rebuild_table(table, headers: list[str], rows: list[list]) -> None:
    """Var olan tabloyu yerinde yeniden doldurur (stil ve konum korunur)."""
    tbl = table._tbl
    for tr in list(tbl.tr_lst):
        tbl.remove(tr)
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for _ in headers:
        gc = OxmlElement("w:gridCol")
        grid.append(gc)
    header_row = table.add_row()
    for j, head in enumerate(headers):
        cell = header_row.cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(head))
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = "" if value is None else str(value)


def table_text(table) -> str:
    return " ".join(cell.text for row in table.rows for cell in row.cells)


def paragraph_before_table(doc: Document, table) -> Paragraph | None:
    prev = table._tbl.getprevious()
    while prev is not None and prev.tag != qn("w:p"):
        prev = prev.getprevious()
    return Paragraph(prev, doc._body) if prev is not None else None


def body_paragraph_from_element(doc: Document, element) -> Paragraph:
    return Paragraph(element, doc._body)


def has_field(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.findall(".//" + qn("w:fldChar"))) or bool(paragraph._p.findall(".//" + qn("w:instrText")))


def norm(text: str) -> str:
    """Türkçe İ/I için güvenli küçük harf karşılaştırması."""
    return text.replace("İ", "i").replace("I", "ı").casefold()


def tr_upper(text: str) -> str:
    return text.replace("i", "İ").replace("ı", "I").upper()


def style_name(paragraph: Paragraph) -> str:
    try:
        return paragraph.style.name or ""
    except Exception:
        return ""


HEADING_NUM_RE = re.compile(r"^\s*(\d+)(?:\.(\d+))?(?:\.(\d+))?\.?\s+(\S.*)$")


def is_heading(paragraph: Paragraph) -> bool:
    text = paragraph.text.strip()
    if not text or len(text) > 140:
        return False
    name = style_name(paragraph).casefold()
    if "heading" in name or "başlık" in name or "baslik" in name:
        return True
    if re.search(r"\.{3,}|…|\t|\s\d{1,3}\s*$", text):
        return False  # içindekiler / dizin satırı
    if HEADING_NUM_RE.match(text) and not text.endswith((".", ";", ",")) and not text.startswith(("Tablo", "Şekil", "Ek ")):
        return True
    return False


def is_styled_heading(paragraph: Paragraph) -> bool:
    name = norm(style_name(paragraph))
    return "heading" in name or "başlık" in name or "baslik" in name


def heading_level(paragraph: Paragraph) -> int:
    name = style_name(paragraph)
    m = re.search(r"(\d)", name)
    if ("heading" in name.casefold() or "başlık" in name.casefold()) and m:
        return int(m.group(1))
    m = HEADING_NUM_RE.match(paragraph.text.strip())
    if m:
        return 1 + sum(1 for g in m.groups()[1:3] if g)
    return 0


def heading_number(paragraph: Paragraph) -> str:
    m = HEADING_NUM_RE.match(paragraph.text.strip())
    if not m:
        return ""
    return ".".join(g for g in m.groups()[:3] if g)


# ----------------------------------------------------------------------------
# metin düzeltmeleri
# ----------------------------------------------------------------------------
# "ticari veri+ek" -> "sentetik çiftlik kaydı/kayıtları+ek" (belirtisiz ad tamlaması çekimi)
TR_SUFFIX = {
    "": "kaydı", "si": "kaydı", "sini": "kaydını", "sinin": "kaydının", "ye": "kaydına", "yi": "kaydını",
    "nin": "kaydının", "den": "kaydından", "de": "kaydında",
    "ler": "kayıtları", "lere": "kayıtlarına", "lerden": "kayıtlarından", "leri": "kayıtlarını",
    "lerin": "kayıtlarının", "lerle": "kayıtlarıyla", "lerinin": "kayıtlarının", "lerini": "kayıtlarını",
    "lerine": "kayıtlarına", "lerde": "kayıtlarında",
}


def _commercial_data(match: re.Match) -> str:
    suffix = match.group(1) or ""
    noun = TR_SUFFIX.get(suffix.casefold())
    if noun is None:
        return match.group(0)
    return f"sentetik çiftlik {noun}"


REGEX_FIXES: list[tuple[re.Pattern, object]] = [
    (re.compile(r"Belirli alanlar veya sektörler hakkında daha detaylı bir analiz isterseniz lütfen bize bildirin\.?\s*"), ""),
    (re.compile(r"Bu yaklaşımlar,[^.]*?daha güvenilir karar destek sistemleri oluşturur\."),
     "Bu yaklaşımlar, yapılandırılmış çiftlik verileri ile kaynak gösterilebilir bilimsel bilginin birlikte "
     "kullanımı yoluyla daha güvenilir karar destek sistemlerinin oluşturulmasına katkı sağlar."),
    (re.compile(r"(Geri\s+(?:Alma|Getirim|Getirimli|Alım)\s*[-–]?\s*(?:Destekli|Artırılmış)\s+Üretim)"), "Geri Getirme Destekli Üretim"),
    (re.compile(r"(Bilgi\s+(?:Alma|Getirme)\s*[-–]?\s*(?:Destekli|Genişletilmiş)\s+Üretim)"), "Geri Getirme Destekli Üretim"),
    (re.compile(r"Geri\s+Getirim\b"), "Geri Getirme"),
    (re.compile(r"\bGeri\s+Alma\b"), "Geri Getirme"),
    (re.compile(r"\bgeri\s+alma\b"), "geri getirme"),
    (re.compile(r"\bAGR\b"), "RAG"),
    (re.compile(r"RAG\s*\((?:İlişkisel Toplu Veritabanı|Yanıt Edinimi ve Üretken Yaklaşım|Alınımı Artırılmış Üretim)\)"), "RAG (Geri Getirme Destekli Üretim)"),
    (re.compile(r"Transformer tabanlı öğrenme yönetim sistemleri"), "Transformer tabanlı büyük dil modelleri"),
    (re.compile(r"\byapay zeka\b"), "yapay zekâ"),
    (re.compile(r"\bYapay zeka\b"), "Yapay zekâ"),
    (re.compile(r"\bYapay Zeka\b"), "Yapay Zekâ"),
    (re.compile(r"\bYAPAY ZEKA\b"), "YAPAY ZEKÂ"),
    (re.compile(r"\bticari veri(ler|lere|lerden|leri|lerin|lerle|lerinin|lerini|lerine|si|sini|sinin|ye|yi|nin|den|de)?\b", re.IGNORECASE), _commercial_data),
    (re.compile(r"\btescilli yapılandırılmış veriler(in|i)?\b"), lambda m: "sentetik yapılandırılmış çiftlik veriler" + (m.group(1) or "")),
    (re.compile(r"\b([Şş])irket(?:e|lere)? özg(?:ü|ü olan)\b"), lambda m: ("İ" if m.group(1) == "Ş" else "i") + "şletmeye özgü (sentetik)"),
    (re.compile(r"\bşirket özel(indeki|i)?\b"), lambda m: "sentetik kayıtlardaki" if m.group(1) == "indeki" else "sentetik çiftlik"),
    (re.compile(r"([a-zçğıöşü\)])\s+(Bu çalışmada)"), r"\1. \2"),
    (re.compile(r"\b(Şekil|Tablo)(\d+)\.(\d+)\.(?=\S)"), r"\1 \2.\3. "),
    (re.compile(r"\b(Şekil|Tablo)(\d+\.\d+)"), r"\1 \2"),
    # v1 betiğinin bastığı eski üretim sayıları (varsa) güncel rapora çekilir
    (re.compile(r"ortalama atıf doğruluğu 0,33, desteksiz iddia oranı ise 0,84"), "ortalama atıf doğruluğu 0,30, desteksiz iddia oranı ise 0,86"),
    (re.compile(r"atıf doğruluğu 0,33, atıf tamlığı 0,16, desteksiz iddia oranı 0,84 ve tercih edilen kaynağa isabet 0,68"),
     "atıf doğruluğu 0,30, atıf tamlığı 0,14, desteksiz iddia oranı 0,86 ve tercih edilen kaynağa isabet 0,68"),
    (re.compile(r"atıf doğruluğu 0,33; desteksiz iddia 0,84"), "atıf doğruluğu 0,30; desteksiz iddia 0,86"),
    (re.compile(r"mean citation correctness was 0\.33 and the unsupported-claim rate was 0\.84"), "mean citation correctness was 0.30 and the unsupported-claim rate was 0.86"),
    (re.compile(r"ortalama süre 40,59 saniye \(medyan 37,35; p95 56,93\)"), "ortalama süre 44,61 saniye (medyan 41,19; p95 60,79)"),
]


def apply_text_fixes(text: str) -> str:
    for pattern, repl in REGEX_FIXES:
        text = pattern.sub(repl, text)  # type: ignore[arg-type]
    return text


def global_text_fixes(doc: Document) -> None:
    n = 0
    for para in doc.paragraphs:
        old = para.text
        new = apply_text_fixes(old)
        if new != old:
            set_para_text(para, new)
            n += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    old = para.text
                    new = apply_text_fixes(old)
                    if new != old:
                        set_para_text(para, new)
                        n += 1
    log("Terim birliği / editoryal kalıntı / sentetik veri ifadeleri", "uygulandı", f"{n} paragraf değişti")


# ----------------------------------------------------------------------------
# çelişen paragraflar
# ----------------------------------------------------------------------------
def replace_contradictions(doc: Document, texts: dict) -> None:
    done = {"4.4": False, "5": False, "sql": False}
    for para in doc.paragraphs:
        t = para.text
        if not done["4.4"] and ("Precision@K" in t or "Precision@5" in t) and ("henüz" in t or "elde edilmemiştir" in t or "dahil edilmemiştir" in t):
            set_para_text(para, texts["sec_4_4_fix"])
            done["4.4"] = True
            continue
        if not done["5"] and re.search(r"[Yy]alnızca LLM", t) and ("FAISS" in t or "hibrit" in t) and ("kesin" in t or "nicel" in t or "tamamlanmamıştır" in t):
            set_para_text(para, texts["sec_5_claim_fix"])
            done["5"] = True
            continue
        if not done["sql"] and "24" in t and re.search(r"23['’]ü", t) and ("SQL" in t or "sorgu" in t):
            set_para_text(para, texts["sec_4_3_sql_fix"])
            done["sql"] = True
            continue
    log("§4.4 'nihai nicel sonuçlar henüz yok' çelişkisi", "uygulandı" if done["4.4"] else "UYGULANAMADI", "paket §5")
    log("§5 'LLM/FAISS/hibrit kesin sayı yok' cümlesi", "uygulandı" if done["5"] else "UYGULANAMADI", "paket §9")
    log("§4.3 SQL 23/24 paragrafı (F-SQL019 açıklaması)", "uygulandı" if done["sql"] else "UYGULANAMADI", "paket §4")


# ----------------------------------------------------------------------------
# bölüm ekleri
# ----------------------------------------------------------------------------
def find_heading(doc: Document, number: str | None = None, contains: str | None = None, level: int | None = None) -> Paragraph | None:
    for styled_only in (True, False):
        for para in doc.paragraphs:
            if not is_heading(para) or (styled_only and not is_styled_heading(para)):
                continue
            if number is not None and heading_number(para) != number:
                continue
            if contains is not None and norm(contains) not in norm(para.text):
                continue
            if level is not None and heading_level(para) not in (level, 0):
                continue
            return para
    return None


def pick_styles(doc: Document) -> dict:
    styles = {"h1": None, "h2": None, "body": None, "caption": None}
    counts: dict[str, int] = {}
    for para in doc.paragraphs:
        if is_styled_heading(para):
            lvl = heading_level(para)
            if lvl == 1 and styles["h1"] is None:
                styles["h1"] = para.style
            elif lvl == 2 and styles["h2"] is None:
                styles["h2"] = para.style
        elif len(para.text) > 200:
            counts[style_name(para)] = counts.get(style_name(para), 0) + 1
        if para.text.strip().startswith(("Tablo ", "Ek Tablo ")) and styles["caption"] is None:
            styles["caption"] = para.style
    if counts:
        best = max(counts, key=counts.get)
        for para in doc.paragraphs:
            if style_name(para) == best:
                styles["body"] = para.style
                break
    if styles["h1"] is None:
        for para in doc.paragraphs:
            if is_heading(para) and heading_level(para) == 1:
                styles["h1"] = para.style
                break
    if styles["h2"] is None:
        styles["h2"] = styles["h1"]
    return styles


def insert_after_heading_block(doc: Document, heading: Paragraph, text: str, styles: dict) -> Paragraph:
    """Başlığın altındaki ilk gövde paragrafından sonra ekler (başlık hemen boşsa başlıktan sonra)."""
    nxt = heading._p.getnext()
    anchor = heading
    while nxt is not None and nxt.tag == qn("w:p"):
        candidate = Paragraph(nxt, heading._parent)
        if is_heading(candidate):
            break
        if candidate.text.strip():
            anchor = candidate
            break
        nxt = nxt.getnext()
    return insert_paragraph_after(anchor, text, styles["body"])


def add_protocol_and_test_set(doc: Document, texts: dict, styles: dict) -> None:
    h37 = find_heading(doc, number="3.7") or find_heading(doc, contains="değerlendirme", level=2)
    if h37 is not None:
        insert_after_heading_block(doc, h37, texts["sec_3_7_protocol"], styles)
        log("§3.7 değerlendirme protokolü / gold etiket açıklaması", "uygulandı", h37.text[:60])
    else:
        log("§3.7 değerlendirme protokolü / gold etiket açıklaması", "UYGULANAMADI", "paket §2")

    anchor = None
    for para in doc.paragraphs:
        if "Macro-F1" in para.text and not is_heading(para) and heading_chapter_of(doc, para) == "4":
            anchor = para
            break
    if anchor is None:
        anchor = find_heading(doc, number="4.2")
    if anchor is not None:
        insert_paragraph_after(anchor, texts["sec_test_set"], styles["body"])
        log("§4.2 held-out küme bağımsızlığı / örtüşme analizi", "uygulandı", anchor.text[:60])
    else:
        log("§4.2 held-out küme bağımsızlığı / örtüşme analizi", "UYGULANAMADI", "paket §3")


def heading_chapter_of(doc: Document, paragraph: Paragraph) -> str:
    chapter = ""
    for para in doc.paragraphs:
        if is_heading(para):
            num = heading_number(para)
            if num:
                chapter = num.split(".")[0]
        if para._p is paragraph._p:
            return chapter
    return chapter


def find_conclusion_heading(doc: Document) -> Paragraph | None:
    """Sonuç bölümünün 1. düzey başlığı (ör. '5. SONUÇ VE ÖNERİLER')."""
    for styled_only in (True, False):
        for para in doc.paragraphs:
            if not is_heading(para) or heading_level(para) not in (1, 0):
                continue
            if styled_only and not is_styled_heading(para):
                continue
            t = norm(para.text.strip())
            if "sonuç" not in t:
                continue
            num = heading_number(para)
            if num.split(".")[0] == "5" or not num:
                return para
    return None


def next_table_number(doc: Document, chapter: str) -> int:
    nums = [0]
    for para in doc.paragraphs:
        m = re.match(rf"^Tablo {chapter}\.(\d+)", para.text.strip())
        if m and not re.search(r"\.{3,}|\s\d{1,3}\s*$", para.text.strip()):
            nums.append(int(m.group(1)))
    return max(nums) + 1


def add_results_sections(doc: Document, data: dict, texts: dict, styles: dict) -> None:
    conclusion = find_conclusion_heading(doc)
    if conclusion is None:
        log("§4.5–4.7 yeni bulgular alt bölümleri", "UYGULANAMADI", "Sonuç başlığı bulunamadı; paket §6–8")
        return
    existing = [heading_number(p) for p in doc.paragraphs if is_heading(p) and heading_number(p).startswith("4.")]
    nums = [int(n.split(".")[1]) for n in existing if n.count(".") == 1 and n.split(".")[1].isdigit()]
    nxt = (max(nums) if nums else 4) + 1
    tno = next_table_number(doc, "4")
    rt = data["retrieval"]
    gn = data["generation"]
    pdd = rt["paired_faiss_minus_hybrid"]

    def f(x, nd=2):
        return ("—" if x is None else f"{x:.{nd}f}").replace(".", ",")

    blocks = []
    blocks.append(("h2", f"4.{nxt}. Eşleştirilmiş Retrieval Farkları ve Birleştirme Ağırlığının Seçimi"))
    blocks.append(("p", texts["sec_4_5_paired"]))
    blocks.append(("caption", f"Tablo 4.{tno}. Soru bazında eşleştirilmiş retrieval farkları (FAISS − hibrit), n = {rt['n']}, 10.000 bootstrap örneği"))
    blocks.append(("table", (
        ["Ölçüt", "FAISS", "Hibrit", "Ortalama fark", "Bootstrap %95 GA", "FAISS daha iyi / eşit / hibrit daha iyi"],
        [[k_tr, f(rt["summary"]["faiss"][k]), f(rt["summary"]["hybrid"][k]), f(pdd[k]["mean_diff"], 3),
          f"{f(pdd[k]['ci95_low'], 3)} … {f(pdd[k]['ci95_high'], 3)}",
          f"{pdd[k]['faiss_better']} / {pdd[k]['ties']} / {pdd[k]['hybrid_better']}"]
         for k, k_tr in (("p_at_k", "P@5"), ("r_at_k", "R@5"), ("mrr", "MRR"), ("ndcg_at_k", "nDCG@5"))],
    )))
    nxt += 1
    gs = gn["summary"]
    blocks.append(("h2", f"4.{nxt}. Üretim Kalitesi ve Hata Tipolojisi"))
    blocks.append(("p", texts["sec_4_6_generation"]))
    blocks.append(("caption", f"Tablo 4.{tno + 1}. Held-out RAG kümesinde otomatik üretim kalitesi ölçütleri (n = {gn['n']}; uzman puanı yoktur)"))
    blocks.append(("table", (
        ["Ölçüt (otomatik)", "Değer"],
        [["Atıf doğruluğu (geçerli K-id / yazılan K-id)", f(gs["mean_citation_correctness"])],
         ["Atıf tamlığı (atıflı cümle / cümle)", f(gs["mean_citation_completeness"])],
         ["Desteksiz iddia oranı (1 − atıf tamlığı)", f(gs["mean_unsupported_claim_rate"])],
         ["Tercih edilen kaynağa isabet", f(gs["mean_preferred_source_hit_rate"])],
         ["Tanı reddi", f"{gs['diagnosis_refusal_ok']}/{gn['n']} maddede kural ihlali yok"]],
    )))
    blocks.append(("caption", f"Tablo 4.{tno + 2}. RAG üretim hata tipolojisi (otomatik sınıflama; bir yanıt birden fazla kategoride sayılabilir)"))
    blocks.append(("table", (["Hata türü", "Sayı"], [[k, v] for k, v in gn["error_categories"].items()])))
    nxt += 1
    la, ll = gn["latency_all"], gn["latency_llm_only"]
    blocks.append(("h2", f"4.{nxt}. Final Yapılandırmada Yanıt Süresi Dağılımı"))
    blocks.append(("p", texts["sec_latency"]))
    blocks.append(("caption", f"Tablo 4.{tno + 3}. Final RAG kanalı yanıt süresi dağılımı (saniye; GPU yok; tek tekrar)"))
    blocks.append(("table", (
        ["Küme", "n", "Ortalama", "Medyan", "SS", "Min", "Maks", "p90", "p95"],
        [["Tüm maddeler", la["n"], f(la["mean"]), f(la["median"]), f(la["sd"]), f(la["min"], 3), f(la["max"]), f(la["p90"]), f(la["p95"])],
         ["LLM çağrısı içeren", ll["n"], f(ll["mean"]), f(ll["median"]), f(ll["sd"]), f(ll["min"]), f(ll["max"]), f(ll["p90"]), f(ll["p95"])]],
    )))

    anchor = insert_paragraph_before(conclusion, "", styles["body"])
    for kind, payload in blocks:
        if kind == "h2":
            anchor = insert_paragraph_after(anchor, payload, styles["h2"])
        elif kind == "p":
            anchor = insert_paragraph_after(anchor, payload, styles["body"])
        elif kind == "caption":
            anchor = insert_paragraph_after(anchor, payload, styles["caption"] or styles["body"])
        elif kind == "table":
            headers, rows = payload
            table = insert_table_after(doc, anchor, headers, rows, table_style=styles.get("table"))
            anchor = insert_paragraph_after(anchor, "", styles["body"])
            table._tbl.addnext(anchor._p)
    log("§4.5–4.7 yeni bulgular alt bölümleri (eşleştirilmiş fark, üretim, gecikme)", "uygulandı", "Sonuç başlığından önce")


def add_discussion_chapter(doc: Document, texts: dict, styles: dict) -> None:
    conclusion = find_conclusion_heading(doc)
    if conclusion is None:
        log("Yeni 5. TARTIŞMA bölümü", "UYGULANAMADI", "Sonuç başlığı bulunamadı; paket §10")
        return
    # Sonuç bölümü başlıklarını 5.x -> 6.x yap
    renumbered = 0
    in_conclusion = False
    for para in doc.paragraphs:
        if para._p is conclusion._p:
            in_conclusion = True
        if in_conclusion and is_heading(para):
            t = para.text.strip()
            m = re.match(r"^5((?:\.\d+)*)(\.?)(\s+.*)$", t)
            if m:
                set_para_text(para, f"6{m.group(1)}{m.group(2)}{m.group(3)}")
                renumbered += 1
            elif heading_level(para) == 1 and heading_number(para) and not heading_number(para).startswith("5"):
                break
    anchor = insert_paragraph_before(conclusion, "", styles["body"])
    for text, kind in texts["discussion"]:
        if kind == "h1":
            anchor = insert_paragraph_after(anchor, text, styles["h1"])
        elif kind == "h2":
            anchor = insert_paragraph_after(anchor, text, styles["h2"])
        elif kind == "table":
            anchor = insert_paragraph_after(anchor, "Tablo 5.1. Araştırma sorularının ölçülen ve ölçülmeyen bileşenleri", styles["caption"] or styles["body"])
            table = insert_table_after(doc, anchor, ["Araştırma sorusu", "Yanıtlanma düzeyi", "Dayanak / kalan iş"], [list(r) for r in texts["rq_table"]], table_style=styles.get("table"))
            anchor = insert_paragraph_after(anchor, "", styles["body"])
            table._tbl.addnext(anchor._p)
        else:
            anchor = insert_paragraph_after(anchor, text, styles["body"])
    log("Yeni 5. TARTIŞMA bölümü (Sonuç 6 olarak yeniden numaralandı)", "uygulandı", f"{renumbered} Sonuç başlığı yeniden numaralandı")


def add_limitations(doc: Document, texts: dict, styles: dict) -> None:
    anchor = None
    for para in doc.paragraphs:
        t = para.text
        if not is_heading(para) and len(t) > 120 and re.search(r"sınırlılık|sınırlamalar", t) and ("sentetik" in t or "100 soruluk" in t or "held-out" in t):
            anchor = para
    if anchor is None:
        for para in doc.paragraphs:
            if is_heading(para) and "sınırl" in norm(para.text):
                anchor = para
                break
    if anchor is None:
        log("Sınırlılıklar ek paragrafı", "UYGULANAMADI", "paket §11")
        return
    insert_paragraph_after(anchor, texts["limitations_additions"], styles["body"])
    log("Sınırlılıklar ek paragrafı", "uygulandı", anchor.text[:60])


# ----------------------------------------------------------------------------
# tablolar
# ----------------------------------------------------------------------------
def find_table(doc: Document, *needles: str):
    for table in doc.tables:
        txt = table_text(table)
        if all(n in txt for n in needles):
            return table
    return None


def replace_sql_appendix(doc: Document, data: dict) -> None:
    table = find_table(doc, "F-SQL001")
    sq = data["sql"]
    headers = ["ID", "Soru", "Çalıştırılan referans sorgu", "Beklenen değer", "Gözlenen değer", "Çalışma durumu", "Karşılaştırma", "Hata türü"]
    rows = [[r["id"], r["question"], r["executed_query"], r["expected_value"], r["observed_value"], r["run_status"], r["comparison_result"], r["error_type"]] for r in sq["rows"]]
    caption = (f"Ek Tablo 1. SQL referans sorgularının PostgreSQL üzerinde çalıştırılma kaydı (n = {sq['n']}; "
               f"{sq['executed_ok']} çalıştı, {sq['failed']} şema hatası: {', '.join(sq['failed_ids'])}). "
               "Bu tablo modelin Text-to-SQL execution accuracy'si değildir.")
    if table is None:
        log("Ek Tablo 1 (24 satır, F-SQL019 dâhil, 8 sütun)", "UYGULANAMADI", "F-SQL001 içeren tablo yok; paket §12")
        return
    rebuild_table(table, headers, rows)
    cap = paragraph_before_table(doc, table)
    if cap is not None and cap.text.strip().startswith(("Ek Tablo 1", "Tablo")):
        set_para_text(cap, caption)
    else:
        prev = table._tbl.getprevious()
        if prev is not None and prev.tag == qn("w:p"):
            insert_paragraph_after(Paragraph(prev, doc._body), caption)
    log("Ek Tablo 1 (24 satır, F-SQL019 dâhil, 8 sütun)", "uygulandı", f"{len(rows)} satır")


def replace_corpus_table(doc: Document, data: dict) -> None:
    table = find_table(doc, "D01", "D16")
    cp = data["corpus"]
    headers = ["ID", "Tam künye", "Yıl", "Tür", "Tema", "Sayfa", "Parça"]
    rows = [[r["id"], r["citation"] + (f" [Teyit: {r['verify']}]" if r["verify"] else ""), r["year"], r["doc_type"], r["theme"], r["pages"] or "—", r["chunks"] or "—"] for r in cp["rows"]]
    if table is None:
        log("Tablo 3.2 corpus manifestosu (tam künye, sayfa, parça)", "UYGULANAMADI", "D01…D16 içeren tablo yok; paket §16")
        return
    rebuild_table(table, headers, rows)
    cap = paragraph_before_table(doc, table)
    if cap is not None and cap.text.strip().startswith("Tablo 3."):
        num = re.match(r"(Tablo \d+\.\d+\.?)", cap.text.strip()).group(1)
        set_para_text(cap, f"{num} Bilgi tabanı corpus manifestosu: {cp['document_count']} belge, {cp['total_pages']} sayfa birimi, {cp['total_chunks']} metin parçası (900 karakter / 150 örtüşme)")
    log("Tablo 3.2 corpus manifestosu (tam künye, sayfa, parça)", "uygulandı", f"{len(rows)} satır")


def extend_config_table(doc: Document, data: dict) -> None:
    table = find_table(doc, "IndexFlatL2") or find_table(doc, "k1") or find_table(doc, "temperature")
    if table is None:
        log("Tablo 3.4 yapılandırma (sürüm, donanım, prompt modu)", "UYGULANAMADI", "paket §17")
        return
    rebuild_table(table, ["Bileşen", "Değer"], [[r["component"], r["value"]] for r in data["config"]])
    log("Tablo 3.4 yapılandırma (sürüm, donanım, prompt modu)", "uygulandı", f"{len(data['config'])} satır")


def add_appendices(doc: Document, data: dict, styles: dict) -> None:
    rg, rt, gn, pr = data["routing"], data["retrieval"], data["generation"], data["prompts"]

    def f(x, nd=2):
        return ("—" if x is None else f"{x:.{nd}f}").replace(".", ",")

    cv = None
    for para in doc.paragraphs:
        if is_heading(para) and "özgeçmiş" in norm(para.text):
            cv = para
            break
    if cv is not None:
        anchor = insert_paragraph_before(cv, "", styles["body"])
    else:
        anchor = doc.add_paragraph("", styles["body"])

    def h1(text):
        nonlocal anchor
        anchor = insert_paragraph_after(anchor, text, styles["h1"])

    def p(text, style=None):
        nonlocal anchor
        anchor = insert_paragraph_after(anchor, text, style or styles["body"])

    def tbl(headers, rows):
        nonlocal anchor
        table = insert_table_after(doc, anchor, headers, rows, table_style=styles.get("table"))
        anchor = insert_paragraph_after(anchor, "", styles["body"])
        table._tbl.addnext(anchor._p)

    h1("EK 2. HELD-OUT TEST KÜMESİ YÖNLENDİRME KAYDI")
    p(f"Ek Tablo 2. Held-out kümesi (n = {rg['n']}) için soru, gold sınıf, tahmin ve yönlendirme mekanizması. "
      f"Çalıştırma zamanı {rg['rows'][0]['run_timestamp_utc']}; yapılandırma özeti {rg['rows'][0]['configuration_hash']}; "
      f"mekanizma dağılımı: {', '.join(f'{k} {v}' for k, v in rg['mechanism_counts'].items())}. "
      f"Pilot kümeyle birebir tekrar: {data['overlap']['exact_duplicates']}; Jaccard ≥ 0,50 soru sayısı: {len(data['overlap']['high_similarity_pairs'])}.",
      styles["caption"] or styles["body"])
    tbl(["ID", "Soru", "Gold", "Tahmin", "Mekanizma"], [[r["query_id"], r["query_text"], r["gold_label"], r["predicted_label"], r["mechanism"]] for r in rg["rows"]])

    h1("EK 3. SORGU BAZLI RETRIEVAL SONUÇLARI")
    p(f"Ek Tablo 3. Held-out RAG soruları (n = {rt['n']}, K = 5); gold etiket belge düzeyinde. RR = reciprocal rank.", styles["caption"] or styles["body"])
    tbl(["ID", "İlgili belge", "BM25 P@5", "BM25 RR", "BM25 nDCG", "FAISS P@5", "FAISS RR", "FAISS nDCG", "Hibrit P@5", "Hibrit RR", "Hibrit nDCG"],
        [[r["query_id"], r["n_relevant_docs"], f(r["bm25_p_at_k"]), f(r["bm25_mrr"]), f(r["bm25_ndcg_at_k"]), f(r["faiss_p_at_k"]), f(r["faiss_mrr"]), f(r["faiss_ndcg_at_k"]), f(r["hybrid_p_at_k"]), f(r["hybrid_mrr"]), f(r["hybrid_ndcg_at_k"])] for r in rt["rows"]])

    h1("EK 4. RAG ÜRETİM DENEYİ SORGU BAZLI SONUÇLAR")
    p(f"Ek Tablo 4. Held-out RAG soruları (n = {gn['n']}); otomatik ölçütler, uzman puanı yoktur.", styles["caption"] or styles["body"])
    tbl(["ID", "Atıf doğ.", "Atıf taml.", "Desteksiz", "Terc. kaynak", "Yazılan K-id", "Cümle", "Kategori", "Süre (s)"],
        [[r["query_id"], f(r["citation_correctness"]), f(r["citation_completeness"]), f(r["unsupported_claim_rate"]), f(r["preferred_source_hit_rate"]), r["cited_ids"] or "—", r["sentence_count"], r["category"] + (" + kesik" if r["truncated"] else ""), f(r["latency_sec"], 1)] for r in gn["rows"]])

    h1("EK 5. PROMPT ŞABLONLARI")
    for key, label in (("router_system", "Yönlendirici sistem istemi (2. aşama, yalnızca kural eşleşmezse)"), ("rag_system", "RAG sistem istemi"), ("rag_human", "RAG kullanıcı istemi"), ("veterinary_disclaimer", "Veteriner uyarı metni")):
        if key in pr:
            p(label, styles["caption"] or styles["body"])
            for line in pr[key].split("\n"):
                p(line if line.strip() else " ")
    log("EK 2–5 (yönlendirme logu, retrieval, üretim, prompt'lar)", "uygulandı", "ÖZGEÇMİŞ'ten önce" if cv is not None else "belge sonu")


# ----------------------------------------------------------------------------
# dizinler ve başlık
# ----------------------------------------------------------------------------
def rebuild_list_index(doc: Document, heading_text: str, caption_prefixes: tuple[str, ...], label: str) -> None:
    heading = None
    for para in doc.paragraphs:
        t = norm(para.text.strip())
        if t in (norm(heading_text), norm(heading_text) + " dizini", norm(heading_text) + " listesi"):
            heading = para
            break
    if heading is None:
        log(f"{label} dizini", "UYGULANAMADI", "başlık bulunamadı")
        return
    entries = []
    nxt = heading._p.getnext()
    while nxt is not None and nxt.tag == qn("w:p"):
        para = Paragraph(nxt, heading._parent)
        text = para.text.strip()
        if text and not text.startswith(caption_prefixes):
            break
        entries.append(para)
        nxt = nxt.getnext()
    if any(has_field(p) for p in entries):
        log(f"{label} dizini", "ATLANDI", "Word alanı (TOC) kullanılıyor; Word'de Ctrl+A → F9 ile güncelleyin")
        return
    captions = []
    seen = set()
    entry_elements = {p._p for p in entries}
    for para in doc.paragraphs:
        t = para.text.strip()
        if para._p in entry_elements or re.search(r"\.{3,}|\s\d{1,3}\s*$", t):
            continue
        if t.startswith(caption_prefixes) and re.match(r"^(Ek )?(Tablo|Şekil) \d", t):
            key = re.match(r"^((Ek )?(Tablo|Şekil) \d+(\.\d+)?[a-z]?)", t).group(1)
            if key in seen:
                continue
            seen.add(key)
            captions.append(t)
    if not captions:
        log(f"{label} dizini", "ATLANDI", "gövdede başlık (caption) bulunamadı; mevcut dizin korundu")
        return
    style = entries[0].style if entries else None
    for para in entries:
        para._p.getparent().remove(para._p)
    anchor = heading
    for cap in captions:
        anchor = insert_paragraph_after(anchor, cap, style)
    log(f"{label} dizini", "uygulandı", f"{len(captions)} girdi (sayfa numaraları Word'de yeniden eklenmeli)")


def retitle(doc: Document, new_title: str) -> None:
    n = 0
    for para in doc.paragraphs:
        t = " ".join(para.text.split())
        if OLD_TITLE_RE.search(t) and len(t) < 220:
            upper = t.isupper()
            set_para_text(para, tr_upper(new_title) if upper else new_title)
            n += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    t = " ".join(para.text.split())
                    if OLD_TITLE_RE.search(t) and len(t) < 220:
                        set_para_text(para, tr_upper(new_title) if t.isupper() else new_title)
                        n += 1
    log("Başlık değişikliği (--retitle)", "uygulandı" if n else "UYGULANAMADI", f"{n} paragraf")


# ----------------------------------------------------------------------------
def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--src", required=True, help="Düzenlenecek tez .docx")
    parser.add_argument("--dst", default=None)
    parser.add_argument("--data", default=str(DEFAULT_DATA))
    parser.add_argument("--retitle", action="store_true")
    parser.add_argument("--no-appendices", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    src = Path(args.src)
    dst = Path(args.dst) if args.dst else src.with_name(src.stem + "_revize_2026-10-10.docx")
    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    texts = data["texts"]

    doc = Document(str(src))
    styles = pick_styles(doc)
    ref_table = find_table(doc, "F-SQL001") or (doc.tables[0] if doc.tables else None)
    styles["table"] = ref_table.style if ref_table is not None else None

    # sıra önemli: önce içerik, sonra dizinler
    replace_contradictions(doc, texts)
    add_protocol_and_test_set(doc, texts, styles)
    add_results_sections(doc, data, texts, styles)
    add_discussion_chapter(doc, texts, styles)
    add_limitations(doc, texts, styles)
    replace_sql_appendix(doc, data)
    replace_corpus_table(doc, data)
    extend_config_table(doc, data)
    if not args.no_appendices:
        add_appendices(doc, data, styles)
    if args.retitle:
        retitle(doc, texts["title_suggestion"])
    global_text_fixes(doc)
    rebuild_list_index(doc, "TABLOLAR", ("Tablo ", "Ek Tablo "), "TABLOLAR")
    rebuild_list_index(doc, "ŞEKİLLER", ("Şekil ",), "ŞEKİLLER")

    print("\nİŞLEM RAPORU")
    print("-" * 78)
    for op, status, detail in LOG:
        print(f"[{status:>13}] {op}" + (f"  —  {detail}" if detail else ""))
    print("-" * 78)
    pending = [op for op, status, _ in LOG if status == "UYGULANAMADI"]
    if pending:
        print("Elle yapılacaklar (paket HTML'den yapıştırın):")
        for op in pending:
            print(f"  • {op}")
    print("Word'de son adım: İÇİNDEKİLER alanını Ctrl+A → F9 ile güncelleyin; başlık değişikliği için --retitle kullanın.")
    if args.dry_run:
        print("(dry-run: dosya kaydedilmedi)")
        return
    doc.save(str(dst))
    print(f"kaydedildi: {dst}")


if __name__ == "__main__":
    main()
